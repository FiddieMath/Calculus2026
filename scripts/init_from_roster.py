#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
从点名册 Excel/CSV 初始化班级：
  - 生成查询码（默认 6 位随机数字；--code sid 表示新学生默认用学号）
  - 生成 classbook.json（教师明文台账，保密）与 keys.csv（发布用，保密）
  - 若提供已有 classbook/keys，则保留旧学生与旧查询码，只为新学生补码

用法示例：
  python scripts/init_from_roster.py ^
      --input "E:\\微积分\\2026秋\\点名册.xlsx" ^
      --outdir "E:\\微积分\\2026秋" ^
      --code random6

运行后请勿把 classbook.json / keys.csv / 查询码*.csv 上传到公开仓库。
"""
import argparse
import csv
import json
import os
import secrets
import sys


def cell_str(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def load_rows(path):
    """读取 xlsx 或 csv/tsv，返回二维字符串列表。"""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, data_only=True)
        ws = wb.worksheets[0]
        rows = [[cell_str(c) for c in row] for row in ws.iter_rows(values_only=True)]
        wb.close()
        return rows
    with open(path, "r", encoding="utf-8-sig", errors="replace") as f:
        text = f.read()
    lines = [ln for ln in text.splitlines() if ln.strip() and not ln.strip().startswith("#")]
    delim = "\t" if "\t" in lines[0] else ","
    return [ln.split(delim) for ln in lines]


def parse_roster(rows):
    header_idx = None
    sid_col = name_col = idx_col = None
    for ri, row in enumerate(rows):
        if "学号" in row and "姓名" in row:
            header_idx = ri
            sid_col = row.index("学号")
            name_col = row.index("姓名")
            if "序号" in row:
                idx_col = row.index("序号")
            break
    if header_idx is None:
        raise SystemExit("未找到包含“学号,姓名”的表头，请检查文件。")

    roster = []
    seen = set()
    for row in rows[header_idx + 1:]:
        if len(row) <= max(sid_col, name_col):
            continue
        sid = cell_str(row[sid_col]) if sid_col < len(row) else ""
        name = cell_str(row[name_col]) if name_col < len(row) else ""
        if not sid or not name:
            continue
        if sid in seen:
            continue
        seen.add(sid)
        origin = cell_str(row[idx_col]) if idx_col is not None and idx_col < len(row) else str(len(roster) + 1)
        roster.append({"origin": origin, "sid": sid, "name": name})
    return roster


def random6(used):
    while True:
        code = f"{secrets.randbelow(1000000):06d}"
        if code not in used:
            return code


def load_json(path):
    if path and os.path.exists(path):
        with open(path, "r", encoding="utf-8-sig") as f:
            return json.load(f)
    return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True, help="点名册 xlsx/csv")
    ap.add_argument("--outdir", required=True, help="输出目录（私密保存）")
    ap.add_argument("--code", default="random6", choices=["random6", "sid"],
                    help="新学生查询码规则：random6=6位随机数字，sid=学号")
    ap.add_argument("--existing-classbook", default=None)
    ap.add_argument("--existing-keys", default=None)
    args = ap.parse_args()

    roster = parse_roster(load_rows(args.input))
    if not roster:
        raise SystemExit("名单为空。")

    os.makedirs(args.outdir, exist_ok=True)
    old_cb = load_json(args.existing_classbook) if args.existing_classbook else None
    old_cb = load_json(os.path.join(args.outdir, "classbook.json")) if not old_cb else old_cb

    old_keys = {}
    if args.existing_keys and os.path.exists(args.existing_keys):
        with open(args.existing_keys, "r", encoding="utf-8-sig") as f:
            for r in csv.reader(f):
                if len(r) >= 3 and r[0] and r[0] != "学号":
                    old_keys[r[0].strip()] = r[2].strip()
    keys_path = os.path.join(args.outdir, "keys.csv")
    if os.path.exists(keys_path) and not old_keys:
        with open(keys_path, "r", encoding="utf-8-sig") as f:
            for r in csv.reader(f):
                if len(r) >= 3 and r[0] and r[0] != "学号":
                    old_keys[r[0].strip()] = r[2].strip()

    used = set(old_keys.values())
    key_map = {}
    old_by_sid = {}
    if old_cb and old_cb.get("students"):
        for s in old_cb["students"]:
            if s.get("sid"):
                old_by_sid[s["sid"]] = s
    new_codes = 0
    for item in roster:
        sid = item["sid"]
        code = old_keys.get(sid)
        if not code:
            code = sid if args.code == "sid" else random6(used)
            used.add(code)
            new_codes += 1
        key_map[sid] = code

    students = []
    for item in roster:
        sid = item["sid"]
        if sid in old_by_sid:
            student = dict(old_by_sid[sid])
            student["name"] = item["name"]
            student["assignments"] = old_by_sid[sid].get("assignments", {}) or {}
            student["_origin"] = item["origin"]
        else:
            student = {
                "sid": sid,
                "name": item["name"],
                "assignments": {},
                "_origin": item["origin"],
            }
        students.append(student)

    students.sort(key=lambda s: int(s["_origin"]) if s["_origin"].isdigit() else 1 << 30)
    for s in students:
        s.pop("_origin", None)

    classbook = {
        "version": 2,
        "course": {
            "courseName": "微积分I（第一学期）",
            "semester": "2026-2027 学年 第1学期",
        },
        "students": students,
    }
    cb_path = os.path.join(args.outdir, "classbook.json")
    with open(cb_path, "w", encoding="utf-8") as f:
        json.dump(classbook, f, ensure_ascii=False, indent=2)

    key_rows = [(s["sid"], s["name"], key_map[s["sid"]]) for s in students]

    with open(os.path.join(args.outdir, "keys.csv"), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["学号", "姓名", "查询码"])
        w.writerows(key_rows)

    handout = []
    for n, (sid, name, code) in enumerate(key_rows, start=1):
        handout.append([str(n), sid, name, code])
    with open(os.path.join(args.outdir, "查询码-2026秋.csv"), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["序号", "学号", "姓名", "查询码"])
        w.writerows(handout)

    print("roster=", len(students), "new_codes=", new_codes)
    print("classbook=", cb_path)
    print("keys=", os.path.join(args.outdir, "keys.csv"))
    print("handout=", os.path.join(args.outdir, "查询码-2026秋.csv"))


if __name__ == "__main__":
    main()
